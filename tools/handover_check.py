#!/usr/bin/env python3
"""Write docs/HANDOVER.xlsx: everything the user still has to do after the skill finished (DOD-05).

Detects what the project contains (flavors, API, Firebase, push, remote config, deep links, analytics providers…),
selects the matching items from handover-catalog.json, pre-fills the status of items it can check
(placeholder google-services.json, empty prod pins/keys, placeholder icon…) and keeps the statuses and notes the
user entered in an existing HANDOVER.xlsx.

Usage:
  generate_handover.py <project-dir> [--out docs/HANDOVER.xlsx]
  tools/handover_check.py            # the same script copied into the project: refreshes the checklist
Needs openpyxl (pip3 install --user openpyxl).
"""
import argparse
import datetime
import glob
import json
import os
import re
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is required: pip3 install --user openpyxl")

KNOWN_PROVIDERS = {"firebase", "adjust", "appsflyer", "mixpanel", "amplitude", "segment"}
STATUSES = ["To do", "Done", "N/A"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(vertical="top", wrap_text=True)


def find_catalog():
    here = os.path.dirname(os.path.abspath(__file__))
    for path in (os.path.join(here, "handover-catalog.json"),
                 os.path.join(here, "..", "references", "handover-catalog.json")):
        if os.path.exists(path):
            return json.load(open(path, encoding="utf-8"))
    sys.exit("handover-catalog.json not found next to the script or in ../references/")


def read(path):
    try:
        return open(path, encoding="utf-8").read()
    except OSError:
        return ""


def gradle_props(root):
    props = {}
    for path in (os.path.join(root, "gradle.properties"), os.path.expanduser("~/.gradle/gradle.properties")):
        for line in read(path).splitlines():
            if "=" in line and not line.lstrip().startswith("#"):
                key, value = line.split("=", 1)
                if value.strip() or key.strip() not in props:
                    props[key.strip()] = value.strip()
    return props


def detect(root):
    app_build = read(os.path.join(root, "app/build.gradle.kts"))
    manifest = read(os.path.join(root, "app/src/main/AndroidManifest.xml"))
    all_builds = "".join(read(p) for p in glob.glob(os.path.join(root, "**/build.gradle.kts"), recursive=True)
                         if "/build/" not in p)
    exists = lambda rel: os.path.exists(os.path.join(root, rel))
    providers = sorted({
        re.sub(r"AnalyticsProvider\.kt$", "", os.path.basename(p)).lower()
        for p in glob.glob(os.path.join(root, "core/analytics/src/main/kotlin/**/providers/*AnalyticsProvider.kt"),
                           recursive=True)
    })
    hosts = sorted(set(re.findall(r'manifestPlaceholders\["deepLinkHost[A-Za-z]*"\]\s*=\s*"([^"]+)"', app_build)))
    features = {
        "always": True,
        "api": exists("core/network"),
        "flavors": "productFlavors" in app_build,
        "firebase": "google.services" in app_build or "firebase" in all_builds,
        "crashlytics": "crashlytics" in app_build,
        "push": exists("core/notifications"),
        "remote_config": exists("core/config"),
        "distribution": exists(".github/workflows/distribute.yml"),
        "deeplinks": "autoVerify" in manifest or bool(hosts),
        "login": exists("feature/auth"),
        "analytics": bool(providers),
    }
    for provider in providers:
        features["analytics:" + (provider if provider in KNOWN_PROVIDERS else "other")] = True
    package = gradle_props(root).get("app.package", "")
    sdks = [p.capitalize() for p in providers] + (["Firebase Crashlytics"] if features["crashlytics"] else []) + \
        (["Firebase Cloud Messaging"] if features["push"] else [])
    return features, {
        "package": package,
        "stage_package": package + ".stage",
        "hosts": ", ".join(hosts) or "your deep-link hosts",
        "prod_host": hosts[-1] if hosts else "<host>",
        "sdk_list": ", ".join(sdks) or "no third-party analytics SDKs",
        "others": [p for p in providers if p not in KNOWN_PROVIDERS],
    }


class Safe(dict):
    def __missing__(self, key):
        return "{" + key + "}"


def fill(text, values):
    return str(text).format_map(Safe(values)) if text else ""


def auto_status(check, root, props):
    """Returns (status, note) for an auto-checkable item, or None when it can't be checked."""
    if not check:
        return None
    if check == "placeholder_icon":
        icon = read(os.path.join(root, "app/src/main/res/drawable/ic_launcher_foreground.xml"))
        return ("To do", "placeholder icon still in place") if "Placeholder launcher glyph" in icon else ("Done", "icon replaced")
    if check == "prod_base_url":
        url = props.get("api.prod.baseUrl") or props.get("api.baseUrl", "")  # api.baseUrl: single-environment app
        bad = not url or re.search(r"localhost|example\.com", url)
        return ("To do", "prod base URL is empty or a placeholder") if bad else ("Done", "set to " + url + " — confirm it is production")
    if check == "prod_pins":
        pins = [p for p in (props.get("api.prod.certPins") or props.get("api.certPins", "")).split(",") if p.strip()]
        return ("Done", "%d pins set — confirm against the prod host" % len(pins)) if pins else ("To do", "prod cert pins are empty")
    if check == "google_services_real":
        path = os.path.join(root, "app/google-services.json")
        if not os.path.exists(path):
            return ("To do", "app/google-services.json is missing")
        text = read(path)
        placeholder = re.search(r"dummy|placeholder|AIza0{35}", text, re.IGNORECASE)
        return ("To do", "placeholder file") if placeholder else ("Done", "real file present")
    if check == "push_token_sink":
        impl = any(re.search(r":\s*PushTokenSink\b", read(p))
                   for p in glob.glob(os.path.join(root, "core/data/**/*.kt"), recursive=True))
        return ("Done", "PushTokenSink bound in :core:data") if impl else ("To do", "no PushTokenSink implementation yet")
    if check.startswith("key:"):
        name = check[4:]
        return ("Done", name + " is set") if props.get(name) else ("To do", name + " is empty")
    return None


def previous_entries(out):
    if not os.path.exists(out):
        return {}
    try:
        ws = load_workbook(out)["Actions"]
    except (KeyError, OSError, ValueError):
        return {}
    header = [c.value for c in ws[1]]
    rows = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        entry = dict(zip(header, row))
        if entry.get("ID"):
            rows[entry["ID"]] = entry
    return rows


def known_exceptions(root):
    spec = read(os.path.join(root, "docs/SPEC.md"))
    items = []
    for section in ("Known exceptions", "Version notes"):
        match = re.search(r"^## %s\s*\n(.*?)(?=^## |\Z)" % section, spec, re.MULTILINE | re.DOTALL)
        if not match:
            continue
        current = None
        for line in match.group(1).splitlines():
            if line.startswith("- "):
                current = [section, line[2:].strip()]
                items.append(current)
            elif current and line.startswith("  "):
                current[1] += " " + line.strip()
    return items


def table(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), len(rows) + 1)


def build(root, out):
    catalog = find_catalog()
    features, values = detect(root)
    props = gradle_props(root)
    previous = previous_entries(out)
    selected = lambda items: [i for i in items if features.get(i["when"])]

    actions = []
    for item in selected(catalog["actions"]):
        providers = values["others"] if item["when"] == "analytics:other" else [None]
        for provider in providers:
            local = dict(values, provider=(provider or "").capitalize())
            item_id = item["id"] + ("-" + provider if provider else "")
            auto = auto_status(item.get("auto"), root, props)
            old = previous.get(item_id, {})
            status, note = (auto if auto else ("To do", ""))
            if old.get("Status") in STATUSES and (not auto or old.get("Status") == "N/A"):
                status = old["Status"]  # manual statuses (and N/A) survive a refresh
            if old.get("Notes") and not auto:
                note = old["Notes"]
            actions.append([item_id, item["area"], fill(item["task"], local), fill(item["why"], local),
                            fill(item["how"], local), item["owner"], "Yes" if item["blocker"] else "No",
                            fill(item["verify"], local), status, ("auto: " + note) if auto else note])

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    ws = wb.create_sheet("Actions")
    table(ws, ["ID", "Area", "Task", "Why", "How", "Owner", "Release blocker", "Verify", "Status", "Notes"],
          actions, (11, 14, 34, 40, 60, 11, 9, 38, 10, 30))
    if actions:
        status_range = "I2:I%d" % (len(actions) + 1)
        validation = DataValidation(type="list", formula1='"%s"' % ",".join(STATUSES), allow_blank=False)
        ws.add_data_validation(validation)
        validation.add(status_range)
        for status, color in (("Done", "C6EFCE"), ("To do", "FFEB9C"), ("N/A", "D9D9D9")):
            ws.conditional_formatting.add(status_range, CellIsRule(operator="equal", formula=['"%s"' % status],
                                                                    fill=PatternFill("solid", fgColor=color)))

    table(wb.create_sheet("Secrets & keys"), ["Name", "What it is", "Where it goes", "Flavor", "How to get it", "Read by"],
          [[fill(s[k], values) for k in ("name", "what", "where", "flavor", "how", "read_by")]
           for s in selected(catalog["secrets"])], (34, 36, 34, 10, 30, 36))
    table(wb.create_sheet("External setup"), ["Service", "Setup", "Action IDs"],
          [[fill(e["service"], values), fill(e["setup"], values), e["ref"]] for e in selected(catalog["external"])],
          (28, 80, 18))
    table(wb.create_sheet("Verify"), ["Check", "Command / steps"],
          [[v["check"], fill(v["command"], values)] for v in selected(catalog["verify"])], (30, 100))
    exceptions = known_exceptions(root)
    table(wb.create_sheet("Known exceptions"), ["Source (SPEC.md)", "Item"],
          exceptions or [["—", "No known exceptions recorded in docs/SPEC.md"]], (20, 110))

    open_items = [a for a in actions if a[8] == "To do"]
    blockers = [a for a in open_items if a[6] == "Yes"]
    summary.append(["Handover checklist — %s" % (values["package"] or os.path.basename(root))])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Generated %s by android-app-builder. Refresh: python3 tools/handover_check.py "
                    "(keeps the statuses and notes you entered)." % datetime.date.today().isoformat()])
    summary.append([])
    summary.append(["Status", "Actions"])
    for status in STATUSES:
        summary.append([status, sum(1 for a in actions if a[8] == status)])
    summary.append(["Release blockers open", len(blockers)])
    summary.append([])
    summary.append(["Open by owner", "Actions"])
    for owner in sorted({a[5] for a in open_items}):
        summary.append([owner, sum(1 for a in open_items if a[5] == owner)])
    summary.append([])
    summary.append(["Release blockers (do these before the first Play upload)", "", "Owner"])
    for blocker in blockers:
        summary.append([blocker[0] + " — " + blocker[2], "", blocker[5]])
    for row in summary.iter_rows():
        if row[0].value in ("Status", "Open by owner") or str(row[0].value).startswith("Release blockers ("):
            for cell in row:
                cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    summary.column_dimensions["A"].width = 70
    summary.column_dimensions["B"].width = 10
    summary.column_dimensions["C"].width = 14

    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    auto_done = sum(1 for a in actions if a[8] == "Done" and str(a[9]).startswith("auto:"))
    print("HANDOVER: %d actions · %d open · %d release blockers open · %d auto-checked done → %s"
          % (len(actions), len(open_items), len(blockers), auto_done, os.path.relpath(out, root)))


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    in_project_tools = os.path.basename(here) == "tools"
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("project", nargs="?", default=os.path.dirname(here) if in_project_tools else None)
    parser.add_argument("--out", default=None)
    args = parser.parse_args()
    if not args.project:
        parser.error("project directory required")
    root = os.path.abspath(args.project)
    build(root, os.path.abspath(args.out or os.path.join(root, "docs", "HANDOVER.xlsx")))


if __name__ == "__main__":
    main()
